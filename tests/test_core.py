import io
from pathlib import Path
import pytest
from app.documents import make_chunks, extract, stored_file
from app.security import password_hash, verify_password

@pytest.mark.parametrize('size,overlap',[(70,10),(200,30),(19,2)])
def test_chunks_preserve_offsets(size,overlap):
    value='Привет, мир! 😊 Проверка фрагментов.\n'*30
    chunks=make_chunks(value,lambda x:len(x)//2,64,target_chars=size,overlap=overlap)
    assert chunks[0].start==0 and chunks[-1].end==len(value)
    assert all(c.content==value[c.start:c.end] for c in chunks)
    assert all(len('passage: '+c.content)//2<=64 for c in chunks)
    assert all(a.start<b.start<=a.end for a,b in zip(chunks,chunks[1:]))

@pytest.mark.parametrize('text',['','   ','\n\n'])
def test_blank_chunks_rejected(text):
    with pytest.raises(ValueError):make_chunks(text,lambda x:1,512)

def test_passwords():
    hashed=password_hash('my-password-123')
    assert hashed!='my-password-123' and verify_password(hashed,'my-password-123')
    assert not verify_password(hashed,'wrong') and not verify_password('invalid','whatever')
    with pytest.raises(ValueError):password_hash('short')

@pytest.mark.parametrize('encoding',['utf-8','cp1251'])
def test_extract_txt(tmp_path,encoding):
    p=tmp_path/'a.txt';p.write_bytes('Привет\r\nмир'.encode(encoding))
    text,pages=extract(p,'txt')
    assert text=='Привет\nмир' and pages is None

def test_extract_docx(tmp_path):
    from docx import Document
    d=Document();d.add_paragraph('Первый абзац');t=d.add_table(rows=1,cols=2)
    t.cell(0,0).text='Лево';t.cell(0,1).text='Право'
    p=tmp_path/'a.docx';d.save(p)
    text,pages=extract(p,'docx')
    assert 'Первый абзац' in text and 'Лево | Право' in text and pages is None

def test_pdf_no_text(tmp_path):
    from pypdf import PdfWriter
    w=PdfWriter();w.add_blank_page(width=100,height=100)
    p=tmp_path/'blank.pdf';w.write(p)
    with pytest.raises(ValueError,match='нет извлекаемого текста'):extract(p,'pdf')

def test_source_path_traversal():
    with pytest.raises(ValueError):stored_file('../../etc/passwd')

@pytest.mark.parametrize('values',[[],[float('nan')],[float('inf')]])
def test_vector_validation(values):
    from app.embeddings import vector_literal
    with pytest.raises(ValueError):vector_literal(values)

def test_pdf_text_and_page_offsets(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer=PdfWriter()
    for content in ['Hello PostgreSQL','Backup and restore']:
        page=writer.add_blank_page(width=300,height=300)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(f'BT /F1 12 Tf 20 200 Td ({content}) Tj ET'.encode())
        page[NameObject('/Contents')]=writer._add_object(stream)
    path=tmp_path/'text.pdf';writer.write(path)
    result,starts=extract(path,'pdf')
    assert 'Hello PostgreSQL' in result and 'Backup and restore' in result
    assert len(starts)==2 and result[starts[1]:].startswith('Backup')
